#!/usr/bin/env python3
"""Inspect, extract, and safely map revised text into a copied DOCX."""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path

COMPLEX_PARTS = {
    "word/comments.xml": "comments",
    "word/footnotes.xml": "footnotes",
    "word/endnotes.xml": "endnotes",
}


def inspect_docx(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        document_xml = archive.read("word/document.xml")
    features = {label: part in names for part, label in COMPLEX_PARTS.items()}
    features["tracked_changes"] = b"<w:ins" in document_xml or b"<w:del" in document_xml
    features["text_boxes"] = b"txbxContent" in document_xml
    present = [name for name, value in features.items() if value]
    return {
        "file": str(path),
        "complex_features": features,
        "requires_warning": bool(present),
        "present": present,
    }


def require_docx():
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for extract/apply") from exc
    return Document


def iter_blocks(document):
    for p_index, paragraph in enumerate(document.paragraphs):
        yield f"p:{p_index}", paragraph, "paragraph"
    for t_index, table in enumerate(document.tables):
        for r_index, row in enumerate(table.rows):
            for c_index, cell in enumerate(row.cells):
                for p_index, paragraph in enumerate(cell.paragraphs):
                    yield f"t:{t_index}:r:{r_index}:c:{c_index}:p:{p_index}", paragraph, "table_cell"


def extract(path: Path) -> dict:
    Document = require_docx()
    document = Document(str(path))
    blocks = []
    for block_id, paragraph, kind in iter_blocks(document):
        blocks.append({
            "id": block_id,
            "kind": kind,
            "style": paragraph.style.name if paragraph.style else None,
            "text": paragraph.text,
        })
    return {"schema_version": "1.0", "source": str(path), "blocks": blocks}


def set_text_preserving_primary_style(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def apply_map(source: Path, mapping_path: Path, output: Path, allow_complex: bool) -> dict:
    if source.resolve() == output.resolve():
        raise ValueError("output must differ from source; overwriting is forbidden")
    inspection = inspect_docx(source)
    if inspection["requires_warning"] and not allow_complex:
        raise RuntimeError(
            "complex DOCX features detected: " + ", ".join(inspection["present"]) +
            "; report them before editing or pass --allow-complex after explicit approval"
        )
    payload = json.loads(mapping_path.read_text(encoding="utf-8-sig"))
    replacements = {item["id"]: item["text"] for item in payload.get("blocks", [])}
    Document = require_docx()
    document = Document(str(source))
    changed = 0
    known_ids = set()
    for block_id, paragraph, _kind in iter_blocks(document):
        known_ids.add(block_id)
        if block_id in replacements and paragraph.text != replacements[block_id]:
            set_text_preserving_primary_style(paragraph, replacements[block_id])
            changed += 1
    unknown = sorted(set(replacements) - known_ids)
    if unknown:
        raise ValueError("mapping contains unknown block ids: " + ", ".join(unknown[:10]))
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return {
        "source": str(source),
        "output": str(output),
        "changed_blocks": changed,
        "primary_paragraph_styles_preserved": True,
        "inline_run_formatting_note": "Changed paragraphs keep the first run style; mixed inline formatting may need visual review.",
        "inspection": inspection,
    }


def write_payload(payload: dict, output: Path | None, pretty: bool) -> None:
    text = json.dumps(payload, ensure_ascii=False, indent=2 if pretty else None)
    if output:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(text + "\n", encoding="utf-8")
    else:
        print(text)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    inspect_parser = sub.add_parser("inspect")
    inspect_parser.add_argument("source", type=Path)
    inspect_parser.add_argument("--output", type=Path)
    inspect_parser.add_argument("--pretty", action="store_true")
    extract_parser = sub.add_parser("extract")
    extract_parser.add_argument("source", type=Path)
    extract_parser.add_argument("--output", type=Path)
    extract_parser.add_argument("--pretty", action="store_true")
    apply_parser = sub.add_parser("apply")
    apply_parser.add_argument("source", type=Path)
    apply_parser.add_argument("mapping", type=Path)
    apply_parser.add_argument("output", type=Path)
    apply_parser.add_argument("--allow-complex", action="store_true")
    apply_parser.add_argument("--report", type=Path)
    apply_parser.add_argument("--pretty", action="store_true")
    args = parser.parse_args()

    try:
        if args.command == "inspect":
            payload = inspect_docx(args.source)
            write_payload(payload, args.output, args.pretty)
        elif args.command == "extract":
            payload = extract(args.source)
            write_payload(payload, args.output, args.pretty)
        else:
            payload = apply_map(args.source, args.mapping, args.output, args.allow_complex)
            write_payload(payload, args.report, args.pretty)
    except (OSError, RuntimeError, ValueError, KeyError, json.JSONDecodeError, zipfile.BadZipFile) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
